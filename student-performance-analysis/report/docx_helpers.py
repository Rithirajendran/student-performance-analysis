"""Small python-docx helper layer used by build_report.py."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

BLUE = RGBColor(0x1F, 0x4E, 0x8C)
GREY = RGBColor(0x52, 0x51, 0x4E)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)


def add_field(run, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); run._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; run._r.append(i)


def new_document():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Heading 1", 17, 16, 8), ("Heading 2", 13.5, 12, 5), ("Heading 3", 11.5, 8, 3)):
        h = d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = BLUE; h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    return d


def para(d, text="", bold=False, italic=False, size=None, align=None, color=None, after=None, style=None, keep=False):
    p = d.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold, italic, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if keep: p.paragraph_format.keep_with_next = True
    return p


def add_runs(p, text, bold=False, italic=False, size=None, color=None, font=None):
    """Supports **bold** inline markup."""
    for i, part in enumerate(text.split("**")):          # **bold**
        for j, sub in enumerate(part.split("*")):        # *italic*
            if not sub: continue
            r = p.add_run(sub); r.bold = bold or (i % 2 == 1); r.italic = italic or (j % 2 == 1)
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if font: r.font.name = font
    return p


def bullets(d, items, style="List Bullet"):
    for n, it in enumerate(items, 1):
        if style == "List Number":      # manual numbering so every list restarts at 1
            p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
            add_runs(p, f"{n}.  " + it)
        else:
            p = d.add_paragraph(style=style); add_runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def table(d, header, rows, widths_cm, font=9, align_right_from=None, caption=None, zebra=True, bold_first_col=False, keep_together=False):
    if caption:
        c = para(d, caption, bold=True, size=10, color=BLUE, after=3, keep=True)
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; set_cell_margins(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.width = Cm(widths_cm[i]); shade(c, "1F4E8C")
        p = c.paragraphs[0]; r = p.add_run(str(h)); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        p.paragraph_format.space_after = Pt(0)
        if align_right_from is not None and i >= align_right_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    repeat_header(t.rows[0])
    for i, w in enumerate(widths_cm): t.columns[i].width = Cm(w)      # grid widths (LibreOffice/Word both honour)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells; no_split(t.rows[-1])
        for i, v in enumerate(row):
            c = cells[i]; c.width = Cm(widths_cm[i])
            if zebra and ri % 2 == 1: shade(c, "F1F5FB")
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            add_runs(p, str(v), bold=(bold_first_col and i == 0), size=font)
            if align_right_from is not None and i >= align_right_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if keep_together:       # keep every row on the same page as the next one (small tables only)
        for r in t.rows[:-1]:
            for c in r.cells:
                for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
    para(d, "", after=4)
    return t


def figure(d, path, width_cm, caption):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = para(d, caption, italic=True, size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    c.paragraph_format.keep_with_next = True
    return c


def code_block(d, text, size=7):
    for line in text.rstrip("\n").split("\n"):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.2)
        r = p.add_run(line if line else " "); r.font.name = "Courier New"; r.font.size = Pt(size)
        r._r.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc_entry(d, text, page, indent_cm=0, bold=False, first=False, last=False, field=None):
    p = d.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Cm(indent_cm); pf.space_after = Pt(2); pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if first and field:
        r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); r._r.append(f)
        r = p.add_run(); i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = field; r._r.append(i)
        r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "separate"); r._r.append(f)
    r = p.add_run(f"{text}\t{page}"); r.bold = bold; r.font.size = Pt(10.5)
    if last and field:
        r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r._r.append(f)
    return p


def footer_with_page(section, left_text):
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]; p.text = ""; p.style = section.part.document.styles["Normal"]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_text + "\tPage "); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    r2 = p.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = GREY; add_field(r2, "PAGE")
